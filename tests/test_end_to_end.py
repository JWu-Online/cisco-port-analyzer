"""
End-to-end test: pretend switches + the real web app.

    python tests/test_end_to_end.py

Starts one pretend switch of each type we use (Cisco IOS and Dell 5548P and
Aruba S3500 over telnet, a Nexus over SSH with its own login), plus one that
rejects the login and one entered with the wrong type. Then drives the app the
way the page does and checks every port's answer, the Excel file, and the
safety checks. No real switches needed.
"""

import io
import os
import sys
import tempfile
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))
os.environ.setdefault("PA_DATA_DIR", tempfile.mkdtemp(prefix="pa-test-"))

from mock_switch import MockCLI, SSHSwitch, TelnetSwitch  # noqa: E402
import samples as S  # noqa: E402
import app as webapp  # noqa: E402


def check(label, condition, detail=""):
    print(("  PASS  " if condition else "  FAIL  ") + label + (f"   ({detail})" if detail and not condition else ""))
    return condition


def wait(client, base):
    started = time.time()
    while client.get("/api/job", **base).json["state"] == "running" and time.time() - started < 180:
        time.sleep(0.3)
    return client.get("/api/job", **base).json


def main():
    ios_cli = MockCLI("C-IDF4-SW1", S.OUTPUTS["cisco_ios"], start_enabled=False, enable_secret="en-secret")
    nxos_cli = MockCLI("LGCC-N7K", S.OUTPUTS["cisco_nxos"])
    dell_cli = MockCLI("B218-5548P", S.OUTPUTS["dell_powerconnect"], start_enabled=False, enable_secret="en-secret")
    aruba_cli = MockCLI("S3500-M205", S.OUTPUTS["aruba_mas"], start_enabled=False, enable_secret="en-secret",
                        prompt_style="aruba")
    bad_cli = MockCLI("B-IDF3-SW3", {})
    wrong_cli = MockCLI("C-IDF9-SW1", S.OUTPUTS["cisco_ios"])

    servers = {
        "ios": TelnetSwitch(ios_cli, "netops", "sw-pass"),
        "nxos": SSHSwitch(nxos_cli, "admin", "core-pass"),
        "dell": TelnetSwitch(dell_cli, "netops", "sw-pass", user_prompt="User Name:"),
        "aruba": TelnetSwitch(aruba_cli, "netops", "sw-pass", user_prompt="User: "),
        "bad": TelnetSwitch(bad_cli, "netops", "another-password"),
        "wrong": TelnetSwitch(wrong_cli, "netops", "sw-pass"),
    }
    addr = {k: f"127.0.0.1:{s.port}" for k, s in servers.items()}
    client = webapp.app.test_client()
    base = {"base_url": "http://127.0.0.1:5000"}
    ok = True

    print("\nSaving the switch list")
    switches = [
        {"address": addr["ios"], "type": "2960x", "protocol": "telnet", "building": "C Building", "location": "C4 IDF"},
        {"address": addr["nxos"], "type": "N7k", "protocol": "ssh", "building": "M Building", "location": "Main MDF"},
        {"address": addr["dell"], "type": "5548P", "protocol": "telnet", "building": "B Building", "location": "B2 IDF"},
        {"address": addr["aruba"], "type": "S3500", "protocol": "telnet", "building": "M Building", "location": "M2 IDF"},
        {"address": addr["bad"], "type": "cisco_ios", "protocol": "telnet", "building": "B Building", "location": "B3 IDF"},
        {"address": addr["wrong"], "type": "dell_powerconnect", "protocol": "telnet", "building": "C Building",
         "location": "C9 IDF"},
    ]
    r = client.post("/api/switches", json={"switches": switches}, **base)
    ok &= check("switch list saved, model names understood as types",
                r.status_code == 200 and [s["type"] for s in r.json["switches"]] ==
                ["cisco_ios", "cisco_nxos", "dell_powerconnect", "aruba_mas", "cisco_ios", "dell_powerconnect"], r.json)
    r = client.post("/api/switches", json={"switches": [{"address": "10.0.0.9", "type": "aruba_aoscx"}]}, **base)
    ok &= check("a type this app can't check is refused", r.status_code == 400)
    r = client.post("/api/settings", json={"default_protocol": "telnet", "parallel": 6, "timeout": 10}, **base)
    ok &= check("settings saved", r.status_code == 200, r.json)

    login = {"username": "netops", "password": "sw-pass", "secret": "en-secret",
             "ssh_different": True, "ssh_username": "admin", "ssh_password": "core-pass"}

    print("\nChecking the run's own checks")
    r = client.post("/api/run", json={"months": 6, "username": "netops"}, **base)
    ok &= check("missing password refused", r.status_code == 400 and "Enter the password." in r.json["problems"])
    r = client.post("/api/run", json={"months": "lots", **login}, **base)
    ok &= check("bad months refused", r.status_code == 400)
    r = client.post("/api/run", json={"months": 6, "scope": "b:Nowhere", **login}, **base)
    ok &= check("unknown building refused", r.status_code == 400 and "No switches match that choice." in r.json["problems"])

    print("\nTest login on one switch")
    r = client.post("/api/test-switch", json={"switch": switches[3], **login}, **base)
    ok &= check("Aruba S3500 test login reads its ports", r.json.get("ok") and r.json.get("ports") == 4
                and r.json.get("name") == "S3500-M205", r.json)
    r = client.post("/api/test-switch", json={"switch": switches[1], **login}, **base)
    ok &= check("Nexus test login uses the SSH login", r.json.get("ok") and r.json.get("ports") == 4, r.json)

    print("\nChecking every switch")
    r = client.post("/api/run", json={"months": 6, "scope": "", **login}, **base)
    ok &= check("run started", r.status_code == 200, r.json)
    snap = wait(client, base)
    for line in snap["lines"]:
        print(f"    {line['t']} [{line['level']}] {line['msg']}")
    ok &= check("run finished", snap["state"] == "done", snap["state"])
    rep = client.get("/api/report", **base).json
    names = {s["address"]: s for s in rep["switches"]}

    print("\nEvery port's answer")
    kind_of = {addr["ios"]: "cisco_ios", addr["nxos"]: "cisco_nxos", addr["dell"]: "dell_powerconnect",
               addr["aruba"]: "aruba_mas", addr["wrong"]: "cisco_ios"}
    for key in ("ios", "nxos", "dell", "aruba"):
        got = {p["port"]: p for p in rep["ports"] if p["address"] == addr[key]}
        for port, (res, vlan, status) in S.EXPECTED[kind_of[addr[key]]].items():
            p = got.get(port, {})
            ok &= check(f"{names[addr[key]]['name']:<11} {port:<9} {res}",
                        (p.get("result"), p.get("vlan"), p.get("status")) == (res, vlan, status),
                        f"got {p.get('result')} {p.get('vlan')!r} {p.get('status')} | {p.get('why')}")
        ok &= check(f"{names[addr[key]]['name']} has no extra ports", set(got) == set(S.EXPECTED[kind_of[addr[key]]]),
                    sorted(set(got) ^ set(S.EXPECTED[kind_of[addr[key]]])))

    print("\nSwitch summaries")
    ok &= check("Cisco summary counts", (names[addr["ios"]]["ports"], names[addr["ios"]]["inactive"]) == (6, 3),
                names[addr["ios"]])
    ok &= check("uptime read on every type", all(names[addr[k]]["uptime_days"] for k in ("ios", "nxos", "dell", "aruba")),
                {k: names[addr[k]]["uptime_days"] for k in ("ios", "nxos", "dell", "aruba")})
    ok &= check("failed login listed with its reason", "login failed" in names[addr["bad"]]["problem"],
                names[addr["bad"]])
    wrong = names[addr["wrong"]]
    ok &= check("wrong type: flagged and what it printed saved",
                wrong["ports"] == 0 and "Is the type right" in wrong["problem"] and wrong["debug_file"]
                and Path(wrong["debug_file"]).exists(), wrong)
    ok &= check("names come from the switch prompts", names[addr["aruba"]]["name"] == "S3500-M205"
                and names[addr["nxos"]]["name"] == "LGCC-N7K")
    ok &= check("closet switch went into enable mode", ios_cli.enabled and dell_cli.enabled and aruba_cli.enabled)
    ok &= check("paging turned off on each type", "terminal length 0" in ios_cli.commands
                and "terminal datadump" in dell_cli.commands and "no paging" in aruba_cli.commands)

    print("\nOne building, and one switch")
    r = client.post("/api/run", json={"months": 6, "scope": "b:M Building", **login}, **base)
    snap = wait(client, base)
    rep_m = client.get("/api/report", **base).json
    ok &= check("M Building: only its 2 switches checked",
                sorted(s["address"] for s in rep_m["switches"]) == sorted([addr["nxos"], addr["aruba"]])
                and snap["lines"][0]["msg"].startswith("Checking the switches in M Building"), snap["lines"][0])
    r = client.post("/api/run", json={"months": 12, "scope": f"s:{addr['ios']}", **login}, **base)
    wait(client, base)
    rep_1 = client.get("/api/report", **base).json
    by_port = {p["port"]: p for p in rep_1["ports"]}
    ok &= check("one switch at 12 months: 7-month-old port is now Active",
                [s["address"] for s in rep_1["switches"]] == [addr["ios"]] and by_port["Gi1/0/2"]["result"] == "ACTIVE"
                and by_port["Gi1/0/4"]["result"] == "INACTIVE", by_port.get("Gi1/0/2"))
    ok &= check("months remembered for next time", client.get("/api/state", **base).json["settings"]["months"] == 12)

    print("\nExcel report")
    from openpyxl import load_workbook
    full = rep["file"]
    inactive_rows = [i for i, p in enumerate(rep["ports"]) if p["result"] == "INACTIVE"]
    r = client.post("/api/excel", json={"file": full, "rows": inactive_rows}, **base)
    wb = load_workbook(io.BytesIO(r.data))
    ports_sheet = list(wb["Ports"].iter_rows(values_only=True))
    ok &= check("Excel has Summary and Ports sheets", wb.sheetnames == ["Summary", "Ports"])
    ok &= check("Ports sheet has exactly the chosen rows", len(ports_sheet) - 1 == len(inactive_rows)
                and all(row[8] == "Inactive" for row in ports_sheet[1:]), len(ports_sheet))
    ok &= check("file named after the months and date", r.headers["X-Filename"].startswith("ports-unused-6mo-"))
    ok &= check("Excel refuses a made-up report name", client.post("/api/excel", json={"file": "../settings.json"},
                                                                     **base).status_code == 404)

    print("\nSafety")
    data_dir = Path(webapp.DATA_DIR)
    saved = "".join(p.read_text(encoding="utf-8", errors="ignore") for p in data_dir.rglob("*") if p.is_file()
                    and p.suffix in (".json", ".csv", ".txt"))
    ok &= check("no password written to disk", not any(pw in saved for pw in ("sw-pass", "core-pass", "en-secret")))
    ok &= check("other host names refused", client.get("/api/state", base_url="http://evil.example:5000").status_code == 403)
    ok &= check("non-JSON posts refused", client.post("/api/run", data={"months": 6}, **base).status_code == 415)
    ok &= check("page loads", client.get("/", **base).status_code == 200)

    for s in servers.values():
        s.stop()
    print("\nALL PASSED" if ok else "\nSOME CHECKS FAILED")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
