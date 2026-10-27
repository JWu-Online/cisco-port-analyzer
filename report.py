"""report.py - the Excel report: a Summary sheet (one row per switch) and a Ports sheet."""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RESULT_LABEL = {"INACTIVE": "Inactive", "UNSURE": "Unsure", "ACTIVE": "Active"}
FILL = {
    "INACTIVE": PatternFill("solid", fgColor="E3E7EB"),
    "UNSURE": PatternFill("solid", fgColor="FCE8CF"),
    "ACTIVE": PatternFill("solid", fgColor="DDF1E3"),
}
HEADER_FILL = PatternFill("solid", fgColor="1C2833")
HEADER_FONT = Font(bold=True, color="FFFFFF")

PORT_COLUMNS = [("Switch", "switch", 18), ("Address", "address", 14), ("Building", "building", 13),
                ("Location", "location", 16), ("Port", "port", 10), ("Description", "description", 26),
                ("Status", "status", 9), ("VLAN", "vlan", 7), ("Result", "result", 9),
                ("Last used", "last_used", 18), ("Days unused", "days_unused", 11),
                ("Why", "why", 52), ("What the switch said", "how", 40)]
SUMMARY_COLUMNS = [("Switch", "name", 18), ("Address", "address", 14), ("Building", "building", 13),
                   ("Location", "location", 16), ("Ports checked", "ports", 12), ("Inactive", "inactive", 9),
                   ("Unsure", "unsure", 9), ("Active", "active", 9), ("Up for (days)", "uptime_days", 12),
                   ("Problem", "problem", 60)]


def _sheet(ws, columns, rows, color_results=False):
    ws.append([c[0] for c in columns])
    for cell in ws[1]:
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    for r in rows:
        values = []
        for _, key, _ in columns:
            v = r.get(key, "")
            if key == "result":
                v = RESULT_LABEL.get(v, v)
            values.append("" if v is None else v)
        ws.append(values)
        if color_results:
            fill = FILL.get(r.get("result"))
            if fill:
                ws.cell(row=ws.max_row, column=[c[1] for c in columns].index("result") + 1).fill = fill
    for n, (_, _, width) in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(n)].width = width
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{ws.max_row}"


def excel_bytes(switches, ports, months, title=""):
    """switches: summary rows; ports: port rows (already filtered to what should be in the report)."""
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    note = f"Ports unused for {months:g}+ months" + (f" - {title}" if title else "")
    summary.append([note])
    summary["A1"].font = Font(bold=True, size=13)
    summary.append([])
    start = summary.max_row + 1
    tmp = wb.create_sheet("tmp")
    _sheet(tmp, SUMMARY_COLUMNS, switches)
    for row in tmp.iter_rows(values_only=True):
        summary.append(list(row))
    for cell in summary[start]:
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
    for n, (_, _, width) in enumerate(SUMMARY_COLUMNS, 1):
        summary.column_dimensions[get_column_letter(n)].width = width
    summary.freeze_panes = f"A{start + 1}"
    wb.remove(tmp)

    _sheet(wb.create_sheet("Ports"), PORT_COLUMNS, ports, color_results=True)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
